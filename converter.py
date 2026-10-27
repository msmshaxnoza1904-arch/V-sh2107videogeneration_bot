import os
import shutil
import uuid
import asyncio
import subprocess
import logging
from PIL import Image

logger = logging.getLogger(__name__)

CONVERT_DIR = "converted"
os.makedirs(CONVERT_DIR, exist_ok=True)


def get_ffmpeg_path() -> str:
    """
    Tizimdan, bin papkasidan yoki WinGet o'rnatgan manzildan ffmpeg ni qidiradi.
    """
    path = shutil.which("ffmpeg")
    if path:
        return path

    # Render yoki Linux muhitida lokal bin papkasi
    local_bin = os.path.join(os.getcwd(), "bin", "ffmpeg")
    if os.path.exists(local_bin):
        return local_bin

    local_app_data = os.environ.get("LOCALAPPDATA", "")
    winget_packages = os.path.join(local_app_data, "Microsoft", "WinGet", "Packages")
    if os.path.exists(winget_packages):
        for root, dirs, files in os.walk(winget_packages):
            if "ffmpeg.exe" in files:
                return os.path.join(root, "ffmpeg.exe")

    return "ffmpeg"


def get_soffice_path() -> str:
    """
    LibreOffice soffice.exe dasturining manzilini topadi.
    """
    path = shutil.which("soffice")
    if path:
        return path

    standard_paths = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"
    ]
    for p in standard_paths:
        if os.path.exists(p):
            return p

    return "soffice"


IMAGE_EXTS = {"jpg", "jpeg", "png", "webp", "bmp", "tiff", "ico"}
VIDEO_EXTS = {"mp4", "avi", "mkv", "mov", "webm", "flv", "3gp"}
AUDIO_EXTS = {"mp3", "wav", "ogg", "m4a", "aac", "flac", "wma", "opus"}
DOC_EXTS = {"pdf", "docx", "doc", "xlsx", "xls", "pptx", "ppt", "txt", "csv"}


def detect_category(file_name: str) -> str | None:
    """
    Fayl nomining kengaytmasiga qarab uning toifasini aniqlaydi.
    """
    _, ext = os.path.splitext(file_name.lower())
    ext = ext.lstrip(".")
    if ext in IMAGE_EXTS:
        return "image"
    if ext in VIDEO_EXTS:
        return "video"
    if ext in AUDIO_EXTS:
        return "audio"
    if ext in DOC_EXTS:
        return "document"
    return None


def get_available_formats(file_name: str) -> tuple[str | None, list[str]]:
    """
    Fayl toifasi va unga mos o'zgartirish formatlarini qaytaradi.
    """
    _, ext = os.path.splitext(file_name.lower())
    ext = ext.lstrip(".")
    cat = detect_category(file_name)

    if cat == "document":
        if ext == "pdf":
            return cat, ["docx", "xlsx", "jpg", "txt"]
        elif ext in ("docx", "doc"):
            return cat, ["pdf", "txt"]
        elif ext in ("xlsx", "xls"):
            return cat, ["pdf", "csv"]
        elif ext in ("pptx", "ppt"):
            return cat, ["pdf"]
        elif ext in ("txt", "csv"):
            return cat, ["pdf", "docx"]

    if cat == "image":
        formats = ["png", "jpg", "webp", "pdf", "bmp"]
        return cat, [f for f in formats if f != ext and not (ext in ("jpeg", "jpg") and f == "jpg")]

    if cat == "video":
        formats = ["mp4", "mp3", "gif", "avi", "mkv", "mov"]
        return cat, [f for f in formats if f != ext]

    if cat == "audio":
        formats = ["mp3", "wav", "ogg", "m4a", "aac", "flac"]
        return cat, [f for f in formats if f != ext]

    return None, []


# =========================================================================
# HUJJATLARNI KONVERTATSIYA QILISH FUNKSIYALARI
# =========================================================================

def _convert_pdf_to_docx(input_path: str, output_path: str):
    """
    PDF faylini erkin tahrirlanadigan Word (DOCX) ga o'giradi.
    - parse_stream_table=False: matnlar sun'iy 40+ ustunli jadvallarga aylanib bloklanib qolishini oldini oladi.
    - Haqiqiy chiziqli jadvallar (parse_lattice_table=True) saqlanadi.
    - Agar pdf2docx da xatolik bo'lsa, zaxira (PyMuPDF -> python-docx) orqali 100% tahrirlanuvchi qilib o'giriladi.
    """
    success = False
    try:
        from pdf2docx import Converter
        cv = Converter(input_path)
        try:
            cv.convert(
                output_path,
                start=0,
                end=None,
                parse_stream_table=False,
                parse_lattice_table=True
            )
            if os.path.exists(output_path) and os.path.getsize(output_path) > 0:
                success = True
        finally:
            cv.close()
    except Exception as e:
        logger.warning(f"pdf2docx orqali o'girishda muammo bo'ldi: {e}")

    # Agar pdf2docx ishlamasa yoki bo'sh fayl chiqarsa, zaxira usuli
    if not success or not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
        logger.info("PyMuPDF va python-docx orqali toza erkin tahrirlanuvchi DOCX yaratilmoqda...")
        _convert_pdf_to_docx_fallback(input_path, output_path)


def _convert_pdf_to_docx_fallback(input_path: str, output_path: str):
    """
    PyMuPDF va python-docx orqali toza, erkin tahrirlanadigan Word hujjati yaratish.
    """
    import pymupdf
    import docx
    from docx.shared import Pt, Cm

    doc = pymupdf.open(input_path)
    word_doc = docx.Document()

    # Standart chet qoldirish (2 sm)
    for section in word_doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    for page_idx, page in enumerate(doc):
        blocks = page.get_text("blocks")
        has_text = False
        for b in blocks:
            if b[6] == 0:  # Matn bloki
                text = b[4].strip()
                if text:
                    p = word_doc.add_paragraph(text)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.line_spacing = 1.15
                    has_text = True

        if not has_text:
            for line in _ocr_page(page).split("\n\n"):
                if line.strip():
                    word_doc.add_paragraph(line.strip())

        if page_idx < len(doc) - 1:
            word_doc.add_page_break()

    word_doc.save(output_path)


def _convert_pdf_to_xlsx(input_path: str, output_path: str):
    """
    PDF faylidagi jadvallar va matnlarni Excel (XLSX) jadvaliga o'giradi.
    """
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = Workbook()
    wb.remove(wb.active)  # Standart bo'sh sahifani o'chirish

    with pdfplumber.open(input_path) as pdf:
        if not pdf.pages:
            raise ValueError("PDF fayli bo'sh!")

        has_data = False
        for page_idx, page in enumerate(pdf.pages):
            sheet_title = f"Sahifa_{page_idx + 1}"[:31]
            ws = wb.create_sheet(title=sheet_title)

            tables = page.extract_tables()
            if tables:
                for table in tables:
                    for row_idx, row in enumerate(table):
                        cleaned_row = [cell if cell is not None else "" for cell in row]
                        ws.append(cleaned_row)
                    ws.append([])  # Jadvallar orasida bo'sh qator
                has_data = True
            else:
                # Agar jadval topilmasa, qatorlardagi matnlarni yozamiz
                text = page.extract_text()
                if text:
                    for line in text.split("\n"):
                        if line.strip():
                            ws.append([line.strip()])
                    has_data = True

    if not has_data:
        # Hech narsa topilmasa, kamida 1 ta sahifa qoldiramiz
        ws = wb.create_sheet(title="Ma'lumot")
        ws.append(["PDF faylida matn yoki jadval ma'lumoti topilmadi."])

    wb.save(output_path)


def _convert_office_to_pdf(input_path: str, target_dir: str) -> str:
    """
    LibreOffice orqali DOCX, DOC, XLSX, XLS, PPTX, PPT larni sifatli PDF ga aylantiradi.
    """
    soffice_bin = get_soffice_path()
    cmd = [
        soffice_bin,
        "--headless",
        "--convert-to",
        "pdf",
        input_path,
        "--outdir",
        target_dir
    ]

    logger.info(f"LibreOffice konvertatsiya qilmoqda: {' '.join(cmd)}")
    result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if result.returncode != 0:
        logger.error(f"LibreOffice xatosi: {result.stderr}")
        raise ValueError("Hujjatni PDF formatiga o'zgartirib bo'lmadi.")

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    expected_output = os.path.join(target_dir, f"{base_name}.pdf")
    if os.path.exists(expected_output):
        return expected_output

    # Agar nom bir oz boshqacha bo'lsa
    candidates = [
        os.path.join(target_dir, f)
        for f in os.listdir(target_dir)
        if f.lower().endswith(".pdf")
    ]
    if candidates:
        return max(candidates, key=os.path.getctime)

    raise FileNotFoundError("PDF fayli yaratilmadi.")


def _convert_with_soffice(input_path: str, target_dir: str, target_format: str) -> str:
    """
    LibreOffice orqali faylni istalgan formatga o'giradi (masalan, DOC -> TXT, XLS -> CSV).
    """
    soffice_bin = get_soffice_path()
    cmd = [soffice_bin, "--headless", "--convert-to", target_format, input_path, "--outdir", target_dir]
    logger.info(f"LibreOffice konvertatsiya qilmoqda: {' '.join(cmd)}")
    result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if result.returncode != 0:
        logger.error(f"LibreOffice xatosi: {result.stderr}")
        raise ValueError("Hujjatni ushbu formatga o'zgartirib bo'lmadi.")

    ext = target_format.split(":")[0]
    candidates = [os.path.join(target_dir, f) for f in os.listdir(target_dir) if f.lower().endswith(f".{ext}")]
    if not candidates:
        raise FileNotFoundError("Natijaviy fayl yaratilmadi.")
    return max(candidates, key=os.path.getctime)


def _convert_docx_to_txt(input_path: str, output_path: str):
    """
    Word (DOCX) faylidagi matn va jadvallarni TXT ga ko'chiradi.
    """
    import docx
    document = docx.Document(input_path)
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def _convert_xlsx_to_csv(input_path: str, output_path: str):
    """
    Excel (XLSX) faylining barcha varaqlarini CSV ga o'giradi (Excel uchun UTF-8 BOM bilan).
    """
    import csv
    from openpyxl import load_workbook

    wb = load_workbook(input_path, read_only=True, data_only=True)
    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for sheet_idx, ws in enumerate(wb.worksheets):
            if len(wb.worksheets) > 1:
                if sheet_idx > 0:
                    writer.writerow([])
                writer.writerow([f"--- {ws.title} ---"])
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if v is None else v for v in row])
    wb.close()


def _read_text_file(input_path: str) -> str:
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            with open(input_path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    return ""


def _convert_text_to_docx(input_path: str, output_path: str):
    """
    TXT yoki CSV faylini tahrirlanadigan Word (DOCX) ga o'giradi. CSV jadval ko'rinishida joylanadi.
    """
    import csv
    import io
    import docx

    content = _read_text_file(input_path)
    word_doc = docx.Document()

    if input_path.lower().endswith(".csv"):
        rows = [r for r in csv.reader(io.StringIO(content)) if r]
        if rows:
            cols = max(len(r) for r in rows)
            table = word_doc.add_table(rows=0, cols=cols)
            table.style = "Table Grid"
            for r in rows:
                cells = table.add_row().cells
                for i, value in enumerate(r):
                    cells[i].text = value
    else:
        for line in content.splitlines():
            word_doc.add_paragraph(line)

    word_doc.save(output_path)


def _ocr_page(page) -> str:
    """
    Skanerlangan (matnsiz) PDF sahifasini Tesseract OCR orqali o'qiydi. Tesseract bo'lmasa bo'sh qaytaradi.
    """
    try:
        import io
        import pytesseract
        pix = page.get_pixmap(dpi=300)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(img, lang="uzb+rus+eng")
    except Exception as e:
        logger.warning(f"OCR ishlamadi: {e}")
        return ""


def _convert_pdf_to_jpg(input_path: str, output_path: str):
    """
    PDF faylining birinchi sahifasini JPG rasmiga aylantiradi.
    """
    import pymupdf
    doc = pymupdf.open(input_path)
    if len(doc) == 0:
        raise ValueError("PDF fayli bo'sh.")
    page = doc.load_page(0)
    pix = page.get_pixmap(dpi=150)
    pix.save(output_path)
    doc.close()


def _convert_pdf_to_txt(input_path: str, output_path: str):
    """
    PDF faylidagi barcha matnlarni TXT ga ko'chiradi.
    """
    import pymupdf
    doc = pymupdf.open(input_path)
    full_text = []
    for page in doc:
        text = page.get_text()
        if not text.strip():
            text = _ocr_page(page)
        full_text.append(text)
    doc.close()

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n\n--- Sahifa ajratgich ---\n\n".join(full_text))


# =========================================================================
# ASOSIY KONVERTATSIYA LOGIKASI
# =========================================================================

def _convert_sync(input_path: str, target_format: str) -> str:
    """
    Faylni belgilangan formatga o'zgartiradi (Sinxron).
    """
    base_name = os.path.splitext(os.path.basename(input_path))[0]
    _, in_ext = os.path.splitext(input_path.lower())
    in_ext = in_ext.lstrip(".")
    out_id = str(uuid.uuid4())[:8]

    target_format = target_format.lower()
    output_filename = f"{base_name}_{out_id}.{target_format}"
    output_path = os.path.join(CONVERT_DIR, output_filename)

    cat = detect_category(input_path)

    # -------------------------------------------------------------
    # 1. Hujjatlar konvertatsiyasi (PDF, Word, Excel)
    # -------------------------------------------------------------
    if in_ext == "pdf":
        if target_format == "docx":
            _convert_pdf_to_docx(input_path, output_path)
            return output_path
        elif target_format == "xlsx":
            _convert_pdf_to_xlsx(input_path, output_path)
            return output_path
        elif target_format == "jpg":
            _convert_pdf_to_jpg(input_path, output_path)
            return output_path
        elif target_format == "txt":
            _convert_pdf_to_txt(input_path, output_path)
            return output_path

    if in_ext in ("docx", "doc", "xlsx", "xls", "pptx", "ppt", "txt", "csv") and target_format == "pdf":
        # Word yoki Excel -> PDF
        temp_dir = os.path.join(CONVERT_DIR, f"odf_{out_id}")
        os.makedirs(temp_dir, exist_ok=True)
        try:
            pdf_path = _convert_office_to_pdf(input_path, temp_dir)
            shutil.move(pdf_path, output_path)
            return output_path
        finally:
            if os.path.exists(temp_dir):
                shutil.rmtree(temp_dir, ignore_errors=True)

    if in_ext == "docx" and target_format == "txt":
        _convert_docx_to_txt(input_path, output_path)
        return output_path

    if in_ext == "xlsx" and target_format == "csv":
        _convert_xlsx_to_csv(input_path, output_path)
        return output_path

    if in_ext in ("txt", "csv") and target_format == "docx":
        _convert_text_to_docx(input_path, output_path)
        return output_path

    if (in_ext == "doc" and target_format == "txt") or (in_ext == "xls" and target_format == "csv"):
        # Eski formatlar uchun LibreOffice
        soffice_fmt = "txt:Text (encoded):UTF8" if target_format == "txt" else "csv:Text - txt - csv (StarCalc):44,34,76"
        temp_dir = os.path.join(CONVERT_DIR, f"odf_{out_id}")
        os.makedirs(temp_dir, exist_ok=True)
        try:
            result_path = _convert_with_soffice(input_path, temp_dir, soffice_fmt)
            shutil.move(result_path, output_path)
            return output_path
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    # -------------------------------------------------------------
    # 2. Rasmlar (Pillow orqali)
    # -------------------------------------------------------------
    if cat == "image" or target_format in ("png", "jpg", "jpeg", "webp", "pdf", "bmp"):
        if in_ext in IMAGE_EXTS and target_format in ("png", "jpg", "jpeg", "webp", "pdf", "bmp"):
            try:
                with Image.open(input_path) as img:
                    target_fmt_upper = target_format.upper()
                    if target_fmt_upper == "JPG":
                        target_fmt_upper = "JPEG"

                    if target_fmt_upper in ("JPEG", "PDF") and img.mode in ("RGBA", "P", "LA"):
                        img = img.convert("RGB")

                    img.save(output_path, target_fmt_upper)
                    return output_path
            except Exception as e:
                logger.error(f"Pillow xatosi: {e}")

    # -------------------------------------------------------------
    # 3. Video va Audio (FFmpeg orqali)
    # -------------------------------------------------------------
    ffmpeg_bin = get_ffmpeg_path()
    cmd = [ffmpeg_bin, "-y", "-i", input_path]

    if target_format == "mp3":
        cmd.extend(["-vn", "-acodec", "libmp3lame", "-q:a", "2"])
    elif target_format == "gif":
        cmd.extend(["-t", "15", "-vf", "fps=10,scale=480:-1:flags=lanczos"])
    elif target_format == "mp4":
        cmd.extend(["-c:v", "libx264", "-preset", "fast", "-c:a", "aac"])
    elif target_format in ("wav", "ogg", "m4a", "aac", "flac"):
        cmd.extend(["-vn"])

    cmd.append(output_path)

    logger.info(f"FFmpeg ishlamoqda: {' '.join(cmd)}")
    result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if result.returncode != 0:
        logger.error(f"FFmpeg xatosi: {result.stderr}")
        raise ValueError("Faylni ushbu formatga o'zgartirib bo'lmadi.")

    if not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
        raise ValueError("Natijaviy fayl yaratilmadi.")

    return output_path


async def convert_file(input_path: str, target_format: str) -> str:
    """
    Asinxron konvertatsiya funksiyasi.
    """
    return await asyncio.to_thread(_convert_sync, input_path, target_format)
